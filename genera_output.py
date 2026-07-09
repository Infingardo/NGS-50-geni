#!/usr/bin/env python3
"""
Rigenera il database delle varianti dentro index.html a partire da NGS_master.xlsx.

L'Excel è la sorgente di verità. index.html è un artefatto: non modificarlo a mano
nella parte `var DB={...};` — le modifiche verrebbero sovrascritte al prossimo run.

    python3 genera_output.py --check    # verifica, non scrive (usare in review)
    python3 genera_output.py            # rigenera index.html

VALIDAZIONE
-----------
Ogni voce viene verificata contro la copertura reale del pannello Diatech
(fonte: "Lista 50 geni.pdf", colonne SNV/InsDel · Fusioni · CNV · MSI, confermata
da pannello_NGS_Diatech_documento_clinico.md §2.3).

Questo controllo esiste perché il tool ha offerto per mesi voci di amplificazione
per FGFR2, FGFR1 e CDK4 — geni che il pannello sequenzia ma sui quali NON chiama
CNV. Una frase di referto può asserire solo ciò che il saggio è in grado di produrre.
"""
import sys, json, re
from pathlib import Path
from collections import OrderedDict, Counter

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("Serve openpyxl:  pip install openpyxl")

ROOT = Path(__file__).parent
XLSX = ROOT / "NGS_master.xlsx"
HTML = ROOT / "index.html"

COLONNE = ["sede", "gene", "v", "reflex", "note", "vafMode", "kind"]

# --- Copertura del pannello: NON modificare senza riscontro sulla lista ufficiale ---
PANEL = {
    "snv": {
        "AKT1","ALK","AR","BRAF","CDK4","CDKN2A","CTNNB1","DDR2","EGFR","ERBB2","ERBB3",
        "ERBB4","ESR1","FGFR1","FGFR2","FGFR3","FGFR4","GNA11","GNAQ","GNAS","HRAS","IDH1",
        "IDH2","KEAP1","KIT","KRAS","MAP2K1","MET","MTOR","NF1","NRAS","NTRK1","NTRK2",
        "NTRK3","PDGFRA","PIK3CA","POLE","PTEN","RAF1","RB1","RET","ROS1","SMAD4","SMO",
        "STK11","TERT","TP53","TSC1",
    },
    "cnv": {"EGFR", "ERBB2", "MET"},
    "fusions": {"ALK","FGFR1","FGFR2","FGFR3","MET","NRG1","NTRK1","NTRK2","NTRK3",
                "PPARG","RET","ROS1"},
}

# A quale colonna del pannello deve appartenere il gene, per ciascun `kind`.
CLASSE = {
    "snv": "snv", "indel": "snv", "splice": "snv", "structural": "snv",
    "del": "snv", "del_hom": "snv",
    "fusion": "fusions",
    "imbalance": "fusions",   # lo sbilanciamento 3'/5' è un surrogato del riarrangiamento
    "cnv": "cnv",
}

# Chiavi che non sono geni del pannello.
NON_GENI = {"MSI", "RAS/RAF"}

# Eccezioni motivate: alterazioni refertabili non coperte dal pannello NGS.
ECCEZIONI = {
    ("PDGFB (DFSP)", "fusion"): "COL1A1::PDGFB nel DFSP — documentata con FISH, non dal pannello",
}


def leggi_master():
    ws = load_workbook(XLSX, read_only=True).active
    righe = list(ws.iter_rows(values_only=True))
    intest = [str(c).strip() if c is not None else "" for c in righe[0]]
    if intest != COLONNE:
        sys.exit(f"Intestazioni inattese in {XLSX.name}:\n  attese: {COLONNE}\n  trovate: {intest}")
    out = []
    for n, r in enumerate(righe[1:], start=2):
        if r[0] is None:
            continue
        out.append({c: ("" if v is None else str(v)) for c, v in zip(COLONNE, r)} | {"_riga": n})
    return out


def valida(righe):
    errori = []
    for r in righe:
        gene, kind, sede = r["gene"], r["kind"], r["sede"]
        if gene in NON_GENI or kind in ("wildtype", "msi"):
            continue
        if (gene, kind) in ECCEZIONI:
            continue
        classe = CLASSE.get(kind)
        if classe is None:
            errori.append(f"riga {r['_riga']}: kind sconosciuto «{kind}» ({gene}/{sede})")
            continue
        if gene not in PANEL[classe]:
            errori.append(
                f"riga {r['_riga']}: {gene} · {sede} · kind={kind} → "
                f"il pannello NON rileva {classe} per {gene}. "
                f"Geni coperti: {', '.join(sorted(PANEL[classe]))}"
            )
        if r["vafMode"] not in ("required", "not_applicable"):
            errori.append(f"riga {r['_riga']}: vafMode «{r['vafMode']}» non valido ({gene}/{sede})")
        if kind in ("cnv", "fusion", "imbalance") and r["vafMode"] != "not_applicable":
            errori.append(f"riga {r['_riga']}: {gene}/{kind} dovrebbe avere vafMode=not_applicable")
    return errori


def costruisci_db(righe):
    db = OrderedDict()
    for r in righe:
        db.setdefault(r["sede"], OrderedDict()).setdefault(r["gene"], []).append(
            {"v": r["v"], "reflex": r["reflex"], "note": r["note"],
             "vafMode": r["vafMode"], "kind": r["kind"]}
        )
    return db


def span_db(src):
    i = src.index("var DB=") + len("var DB=")
    d = 0
    for j in range(i, len(src)):
        if src[j] == "{":
            d += 1
        elif src[j] == "}":
            d -= 1
            if d == 0:
                return i, j + 1
    sys.exit("`var DB={...}` non trovato in index.html")


def main():
    check = "--check" in sys.argv
    righe = leggi_master()

    errori = valida(righe)
    if errori:
        print(f"✗ {len(errori)} voci non compatibili con la copertura del pannello:\n")
        for e in errori:
            print("   " + e)
        print("\nNessun file scritto.")
        return 1

    db = costruisci_db(righe)
    nuovo = json.dumps(db, ensure_ascii=False)

    src = HTML.read_text(encoding="utf8")
    a, b = span_db(src)
    identico = src[a:b] == nuovo

    kinds = Counter(r["kind"] for r in righe)
    print(f"✓ {len(righe)} voci · {len(db)} sedi · validate contro il pannello")
    print(f"  kind: {dict(sorted(kinds.items()))}")

    if check:
        print("✓ index.html è allineato all'Excel" if identico
              else "✗ index.html DIVERGE dall'Excel — esegui senza --check per rigenerarlo")
        return 0 if identico else 1

    if identico:
        print("  index.html già allineato, nessuna scrittura.")
        return 0
    HTML.write_text(src[:a] + nuovo + src[b:], encoding="utf8")
    print(f"✓ index.html rigenerato")
    return 0


if __name__ == "__main__":
    sys.exit(main())
